"""Statement export, pure: konsol/close/statement_export_model.py (konsol#305
story 8.5; decision #305-W5-3, Deepak Pai, 6 Oct 2026: a downloaded .xlsx
carrying the legend, the Provisional/Signed label, the commentary and the
drill rows; rejected: an Excel add-in function).

``workbook(payload, drills)`` turns ``statement_api.get_statement``'s ``ok``
payload and one ``drill_model.drill`` result per statement heading into the
bytes of an .xlsx with three sheets:

- **Statement**: the header, the sign-off label, the legend and the notes
  (declared-accounts gap, comparison note, entities not included), then
  each section with the Numbers screen's own columns (``COLUMNS``),
  heading/"of which"/net result/not-in-chart/residual rows, and each
  heading's commentary text and byline.
- **Drill**: every heading's drill rows in statement order, one line per
  account (a row with no accounts — CTA, outside scope, current-year
  result, rounding — is one line), then the heading's total.
- **Journals**: the top-side journals behind each heading's "Top-side
  journals" drill row, with their basis (posted in this period).

Amounts are written as numbers with a bracket format, the payload's own
sign — never re-derived. A missing amount raises (never written as 0); a
comparison the statement declared not loaded is the text ``NOT_LOADED``,
as on the screen. A non-ok payload raises with its own message: the caller
refuses, it never sends an empty file.

The label texts, the section columns and ``NOT_LOADED`` are shared with
close-ui's ``numbers.js``; ``test_close_statement_export_model.py`` reads
that file and fails when the two drift.

Pure: no frappe and no konsol imports (openpyxl only).
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

PL = "Profit and Loss"
BS = "Balance Sheet"

#: numbers.js ``COLUMNS``.
COLUMNS = {
    PL: ["This period", "Comparison", "Variance", "Year to date"],
    BS: ["This period", "Comparison", "Variance"],
}
_FIELD = {"This period": "current", "Comparison": "comparison",
          "Variance": "variance", "Year to date": "ytd"}
_COMPARISON_FIELDS = ("comparison", "variance")

#: numbers.js ``SIGNOFF_LABEL`` texts.
SIGNOFF_LABEL = {
    "signed": "Signed",
    "provisional": "Provisional",
    "resign_needed": "Re-sign needed — numbers changed after signing",
}

#: numbers.js ``comparisonCell``'s text for a comparison the statement
#: declared not loaded.
NOT_LOADED = "not loaded"

#: Negatives in brackets, as numbers.js ``amountText`` shows them.
AMOUNT_FORMAT = "#,##0.00;(#,##0.00);0.00"

_INDENT = "    "
_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=14)

DRILL_HEADER = ["Section", "Heading", "Heading name", "Layer", "Entity",
                "Account", "Account name", "Amount"]
JOURNAL_HEADER = ["Heading", "Heading name", "Journal", "Description", "Amount",
                  "Posted by", "Approved by", "Basis"]
_BASIS_NOTE = ("Each heading's amount is its This period figure: cumulative "
               "for the Balance Sheet, the period alone for Profit and Loss.")


def _amount(value, what):
    if value is None:
        raise ValueError(f"{what}: missing amount, never written as 0")
    value = float(value)
    return 0.0 if value == 0 else value


def _cell(line, field, comparison_note, what):
    value = line.get(field)
    if field in _COMPARISON_FIELDS and comparison_note and value is None:
        return NOT_LOADED
    return _amount(value, f"{what} {field}")


def _header_text(payload):
    parts = [payload["period"].get("code"), payload.get("consolidation_group"),
             payload.get("reporting_currency")]
    return " · ".join(["Numbers"] + [p for p in parts if p not in (None, "")])


def _label(signoff):
    state = (signoff or {}).get("state")
    if state not in SIGNOFF_LABEL:
        raise ValueError(f"unknown sign-off state: {state}")
    return SIGNOFF_LABEL[state]


def _not_included_text(not_included):
    """numbers.js ``notIncludedView``'s sentence with every code (a file has
    no "show all" toggle). ``None`` when nothing is missing."""
    if not not_included or not not_included.get("count"):
        return None
    count = not_included["count"]
    text = "%d %s in scope %s not in these numbers" % (
        count, "entity" if count == 1 else "entities", "is" if count == 1 else "are")
    codes = not_included.get("entities") or []
    if codes:
        text += ": " + ", ".join(codes)
    if not_included.get("hidden"):
        text += ", and %d outside your scope" % not_included["hidden"]
    return text


def _byline(entry):
    """``<by> · <YYYY-MM-DD HH:MM ±HH:MM>``: the commentary's author and its
    zoned time (``statement_api._commentary``'s ``at``, already in the site's
    zone), written out in full — a file has no "now" to be relative to."""
    at = datetime.fromisoformat(entry["at"])
    stamp = at.strftime("%Y-%m-%d %H:%M")
    offset = at.strftime("%z")
    if offset:
        stamp += " %s:%s" % (offset[:3], offset[3:])
    return "%s · %s" % (entry.get("by"), stamp)


class _Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.row = 1

    def write(self, values, font=None, amount_cols=()):
        for col, value in enumerate(values, start=1):
            if value is None:
                continue
            cell = self.ws.cell(row=self.row, column=col, value=value)
            if font is not None:
                cell.font = font
            if col in amount_cols and isinstance(value, float):
                cell.number_format = AMOUNT_FORMAT
                cell.alignment = Alignment(horizontal="right")
        self.row += 1

    def blank(self):
        self.row += 1


def _amount_cols(start, count):
    return set(range(start, start + count))


def _section_rows(section, commentary, comparison_note):
    """``[(values, font)]`` for one section's table: header row, then each
    line (and a heading's "of which" rows, and a residual's explanation)."""
    name = section["section"]
    if name not in COLUMNS:
        raise ValueError(f"unknown statement section: {name}")
    columns = COLUMNS[name]
    out = [(["Line", "Code"] + columns + ["Commentary", "By"], _BOLD)]

    def amounts(line, what, fields=None):
        return [
            (_cell(line, _FIELD[c], comparison_note, what)
             if fields is None or _FIELD[c] in fields else None)
            for c in columns
        ]

    for line in section["lines"]:
        kind = line.get("kind")
        if kind == "heading":
            code = line["heading"]
            entry = commentary.get(code)
            note = [entry["text"], _byline(entry)] if entry and entry.get("text") else [None, None]
            out.append(([line.get("heading_name"), code] + amounts(line, code) + note, None))
            for include in line.get("includes") or []:
                out.append(([_INDENT + include["label"], None]
                            + amounts(include, include["label"], ("current", "comparison")), None))
        elif kind in ("no_heading", "net_result"):
            font = _BOLD if kind == "net_result" else None
            out.append(([line["label"], None] + amounts(line, line["label"]), font))
        elif kind == "not_in_chart":
            label = "%s (%s)" % (line["label"], ", ".join(line.get("codes") or []))
            out.append(([label, None] + amounts(line, line["label"]), None))
        elif kind == "residual":
            out.append(([line["label"], None] + amounts(line, line["label"], ("current",)), _BOLD))
            if line.get("current"):
                for item in line.get("explained") or []:
                    out.append(([_INDENT + item["label"], None,
                                 _amount(item["amount"], item["label"])], None))
                out.append(([_INDENT + "Unexplained", None,
                             _amount(line.get("unexplained"), "Unexplained")], None))
        else:
            raise ValueError(f"unknown statement line kind: {kind}")
    return out


def _statement_sheet(ws, payload):
    sheet = _Sheet(ws)
    statement = payload["statement"]
    sheet.write([_header_text(payload)], _TITLE)
    sheet.write([_label(payload.get("signoff"))], _BOLD)
    sheet.write([statement["legend"]])
    gap = payload.get("gap")
    comparison_note = statement["periods"].get("comparison_note")
    for note in ((gap or {}).get("message"), comparison_note,
                 _not_included_text(payload.get("not_included"))):
        if note:
            sheet.write([note])
    sheet.blank()

    commentary = payload.get("commentary") or {}
    for section in statement["sections"]:
        sheet.write([section["section"]], _BOLD)
        cols = _amount_cols(3, len(COLUMNS.get(section["section"], ())))
        for values, font in _section_rows(section, commentary, comparison_note):
            sheet.write(values, font, cols)
        sheet.blank()

    ws.column_dimensions["A"].width = 44
    for letter in "CDEF":
        ws.column_dimensions[letter].width = 16
    ws.column_dimensions["G"].width = 48
    ws.column_dimensions["H"].width = 36


def _headings_in_order(statement):
    return [line["heading"] for section in statement["sections"]
            for line in section["lines"] if line.get("kind") == "heading"]


def _drill_sheets(drill_ws, journal_ws, statement, drills):
    drill_sheet, journal_sheet = _Sheet(drill_ws), _Sheet(journal_ws)
    headings = _headings_in_order(statement)
    notes = {d.get("dimensions_note") for d in drills.values() if d.get("dimensions_note")}
    drill_sheet.write([" ".join(sorted(notes) + [_BASIS_NOTE])])
    drill_sheet.write(DRILL_HEADER, _BOLD)
    journal_sheet.write(["Top-side journals posted in this period, by heading."])
    journal_sheet.write(JOURNAL_HEADER, _BOLD)
    amount_col = {DRILL_HEADER.index("Amount") + 1}
    journal_amount_col = {JOURNAL_HEADER.index("Amount") + 1}

    for code in headings:
        drill = drills.get(code)
        if drill is None:
            raise ValueError(f"{code} has no drill")
        name, section = drill.get("heading_name"), drill.get("section")
        for row in drill["rows"]:
            layer = row.get("label") or "Entity"
            entity = row.get("entity")
            what = f"{code} {layer}"
            if row.get("accounts"):
                for account in row["accounts"]:
                    drill_sheet.write(
                        [section, code, name, layer, entity, account["main_account"],
                         account.get("account_name"), _amount(account["amount"], what)],
                        amount_cols=amount_col)
            else:
                drill_sheet.write(
                    [section, code, name, layer, entity, None, None,
                     _amount(row["amount"], what)], amount_cols=amount_col)
            for journal in row.get("journals") or []:
                journal_sheet.write(
                    [code, name, journal["journal_id"], journal.get("description"),
                     _amount(journal["amount"], f"{code} {journal['journal_id']}"),
                     journal.get("posted_by"), journal.get("approved_by"),
                     row.get("journals_basis")], amount_cols=journal_amount_col)
        drill_sheet.write([section, code, name, "Total", None, None, None,
                           _amount(drill["total"], f"{code} total")],
                          _BOLD, amount_col)

    for ws, widths in ((drill_ws, (16, 10, 28, 30, 12, 12, 30, 16)),
                       (journal_ws, (10, 28, 16, 40, 16, 28, 28, 40))):
        for i, width in enumerate(widths):
            ws.column_dimensions[chr(ord("A") + i)].width = width


def workbook(payload, drills):
    """The .xlsx bytes for an ``ok`` statement payload and its drills
    (``{heading: drill_model.drill(...)}``, one per statement heading).
    Raises ``ValueError`` on a non-ok payload (its own message), a missing
    drill, a missing amount, or a kind/state this module does not know."""
    if payload.get("state") != "ok" or not payload.get("statement"):
        raise ValueError(payload.get("message") or "The statement is not ready to export.")
    book = Workbook()
    statement_ws = book.active
    statement_ws.title = "Statement"
    _statement_sheet(statement_ws, payload)
    _drill_sheets(book.create_sheet("Drill"), book.create_sheet("Journals"),
                  payload["statement"], drills)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
