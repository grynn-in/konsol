"""Statement export model, pure: konsol/close/statement_export_model.py
(konsol#305 story 8.5; decision #305-W5-3, 6 Oct: a downloaded .xlsx with
the legend, the Provisional/Signed label, the commentary and a drill sheet;
an Excel add-in function was rejected).

Every expectation is read from the REAL golden payloads the statement
endpoints produce (``close_statement_payload.json`` from ``get_statement``,
``close_drill_payload.json`` from ``get_drill``), never a hand-built dict.
The workbook is opened back with openpyxl and its cells are compared with
those payloads.

The label texts, the section columns and the "not loaded" wording are
duplicated from close-ui's numbers.js (two languages, one screen and one
file). The drift tests below read numbers.js's source and fail when the two
disagree.
"""
import ast
import copy
import importlib.util
import io
import json
import os
import re

import pytest
from openpyxl import load_workbook

APP_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPO_DIR = os.path.dirname(APP_DIR)
FIXTURE = os.path.join(APP_DIR, "tests", "fixtures", "close_statement_payload.json")
FIXTURE_DRILL = os.path.join(APP_DIR, "tests", "fixtures", "close_drill_payload.json")
NUMBERS_JS = os.path.join(REPO_DIR, "close-ui", "src", "numbers.js")

_PATH = os.path.join(APP_DIR, "close", "statement_export_model.py")
_spec = importlib.util.spec_from_file_location("statement_export_model_under_test", _PATH)
M = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(M)


def _payload():
    with open(FIXTURE) as fh:
        return json.load(fh)


def _drills():
    """The golden drill (heading "1"), keyed by heading as the API passes it,
    plus an empty drill for every other statement heading (the model needs
    one per heading; their rows are not under test here — the API test
    drives the real drill_model for all four)."""
    with open(FIXTURE_DRILL) as fh:
        drill = json.load(fh)["drill"]
    out = {drill["heading"]: drill}
    for code, name, section in (("4", "COST OF SALES", "Profit and Loss"),
                                ("2", "LIABILITIES", "Balance Sheet"),
                                ("3", "EQUITY", "Balance Sheet")):
        out[code] = {"heading": code, "heading_name": name, "section": section,
                     "rows": [], "total": 0.0, "dimensions_note": drill["dimensions_note"]}
    return out


def _book(payload=None, drills=None):
    data = M.workbook(payload or _payload(), _drills() if drills is None else drills)
    assert isinstance(data, bytes)
    return load_workbook(io.BytesIO(data))


def _rows(sheet):
    return [list(r) for r in sheet.iter_rows(values_only=True)]


def _row_starting(rows, first, second=None):
    for row in rows:
        if row[0] == first and (second is None or row[1] == second):
            return row
    raise AssertionError("no row starting %r %r in %r" % (first, second, rows))


def _section_block(rows, section):
    """The rows from ``section``'s title row to the next blank row."""
    start = next(i for i, r in enumerate(rows) if r[0] == section)
    out = []
    for row in rows[start + 1:]:
        if all(v is None for v in row):
            break
        out.append(row)
    return out


# --- sheet 1: the statement ---------------------------------------------------

def test_sheets_are_statement_drill_and_journals():
    book = _book()
    assert book.sheetnames == ["Statement", "Drill", "Journals"]


def test_header_label_and_legend_come_from_the_payload():
    payload = _payload()
    rows = _rows(_book()["Statement"])
    assert rows[0][0] == "Numbers · FY2025P07 · G1 · USD"
    assert rows[1][0] == "Provisional"
    assert rows[2][0] == payload["statement"]["legend"]


@pytest.mark.parametrize("state,text", [
    ("signed", "Signed"),
    ("provisional", "Provisional"),
    ("resign_needed", "Re-sign needed — numbers changed after signing"),
])
def test_label_follows_the_signoff_state(state, text):
    payload = _payload()
    payload["signoff"] = {"state": state, "run": "AR-1"}
    assert _rows(_book(payload)["Statement"])[1][0] == text


def test_an_unknown_signoff_state_raises():
    payload = _payload()
    payload["signoff"] = {"state": "maybe", "run": None}
    with pytest.raises(ValueError, match="maybe"):
        M.workbook(payload, _drills())


def test_notes_carry_the_comparison_note_and_not_included_text():
    rows = _rows(_book()["Statement"])
    firsts = [r[0] for r in rows]
    assert "No rows in the warehouse for FY2025 P06" in firsts
    assert "1 entity in scope is not in these numbers: ZZB" in firsts


def test_the_declared_accounts_gap_is_written_when_present():
    payload = _payload()
    payload["gap"] = {"message": "Declare the CTA account in Close Settings."}
    firsts = [r[0] for r in _rows(_book(payload)["Statement"])]
    assert "Declare the CTA account in Close Settings." in firsts


def test_each_section_has_the_screens_columns():
    rows = _rows(_book()["Statement"])
    pl = _section_block(rows, "Profit and Loss")
    bs = _section_block(rows, "Balance Sheet")
    assert pl[0][:8] == ["Line", "Code", "This period", "Comparison", "Variance",
                         "Year to date", "Commentary", "By"]
    assert bs[0][:7] == ["Line", "Code", "This period", "Comparison", "Variance",
                         "Commentary", "By"]


def test_every_heading_amount_equals_the_payload_line():
    payload = _payload()
    rows = _rows(_book()["Statement"])
    for section in payload["statement"]["sections"]:
        block = _section_block(rows, section["section"])
        for line in section["lines"]:
            if line["kind"] != "heading":
                continue
            row = _row_starting(block, line["heading_name"], line["heading"])
            assert row[2] == line["current"], line["heading"]
            # comparison_note is set and comparison is None: "not loaded",
            # never 0 (numbers.js comparisonCell).
            assert row[3] == "not loaded"
            assert row[4] == "not loaded"
            if section["section"] == "Profit and Loss":
                assert row[5] == line["ytd"]


def test_amounts_are_numbers_with_a_bracket_format():
    sheet = _book()["Statement"]
    for row in sheet.iter_rows():
        if row[0].value == "ASSETS":
            assert isinstance(row[2].value, float)
            assert row[2].number_format == M.AMOUNT_FORMAT
            assert "(" in M.AMOUNT_FORMAT
            return
    raise AssertionError("no ASSETS row")


def test_includes_are_indented_rows_after_their_heading():
    rows = _rows(_book()["Statement"])
    bs = _section_block(rows, "Balance Sheet")
    labels = [r[0] for r in bs]
    i = labels.index("EQUITY")
    assert labels[i + 1] == "    of which currency translation (CTA)"
    assert bs[i + 1][2] == 5.07
    assert labels[i + 2] == "    of which current-year result"
    assert bs[i + 2][2] == 241.43


def test_net_result_and_residual_rows_are_written():
    rows = _rows(_book()["Statement"])
    pl = _section_block(rows, "Profit and Loss")
    net = _row_starting(pl, "Net result")
    assert net[2] == 241.43 and net[5] == 241.43
    bs = _section_block(rows, "Balance Sheet")
    residual = _row_starting(bs, "Unmatched residual")
    assert residual[2] == 0  # -0.0 written as 0
    assert residual[3] is None  # the residual line carries only "current"


def test_a_non_zero_residual_writes_its_explained_and_unexplained_rows():
    payload = _payload()
    residual = payload["statement"]["sections"][1]["lines"][-1]
    residual.update({"current": 12.5, "explained": [
        {"label": "CTA not placed (declare the CTA account)", "amount": 10.0}],
        "unexplained": 2.5})
    bs = _section_block(_rows(_book(payload)["Statement"]), "Balance Sheet")
    assert _row_starting(bs, "    CTA not placed (declare the CTA account)")[2] == 10.0
    assert _row_starting(bs, "    Unexplained")[2] == 2.5


def test_heading_commentary_text_and_byline_are_on_the_heading_row():
    payload = _payload()
    entry = payload["commentary"]["4"]
    pl = _section_block(_rows(_book()["Statement"]), "Profit and Loss")
    row = _row_starting(pl, "COST OF SALES", "4")
    assert row[6] == entry["text"]
    assert row[7] == "Zz Analyst · 2025-08-01 10:00 +01:00"
    bs = _section_block(_rows(_book()["Statement"]), "Balance Sheet")
    assert _row_starting(bs, "ASSETS", "1")[5] is None


def test_a_cleared_commentary_is_not_written():
    payload = _payload()
    payload["commentary"]["4"]["text"] = ""
    pl = _section_block(_rows(_book(payload)["Statement"]), "Profit and Loss")
    row = _row_starting(pl, "COST OF SALES", "4")
    assert row[6] is None and row[7] is None


def test_a_null_amount_without_a_comparison_note_raises_never_zero():
    payload = _payload()
    payload["statement"]["periods"]["comparison_note"] = None
    with pytest.raises(ValueError, match="never written as 0"):
        M.workbook(payload, _drills())


def test_a_non_ok_payload_is_refused_with_its_own_message():
    payload = _payload()
    payload.update({"state": "setup_gap", "message": "statement_heading_side_undeclared: 2",
                    "statement": None})
    with pytest.raises(ValueError, match="statement_heading_side_undeclared: 2"):
        M.workbook(payload, {})


def test_an_unknown_line_kind_raises():
    payload = _payload()
    payload["statement"]["sections"][0]["lines"].append({"kind": "mystery", "current": 1.0})
    with pytest.raises(ValueError, match="mystery"):
        M.workbook(payload, _drills())


# --- sheet 2: the drill -------------------------------------------------------

def test_drill_sheet_has_one_line_per_account_and_sums_to_the_total():
    drill = _drills()["1"]
    rows = _rows(_book()["Drill"])
    header = rows[1]
    assert header == ["Section", "Heading", "Heading name", "Layer", "Entity",
                      "Account", "Account name", "Amount"]
    assert drill["dimensions_note"] in rows[0][0]
    body = [r for r in rows[2:] if r[1] == "1" and r[3] != "Total"]
    expected = []
    for d in drill["rows"]:
        layer = d["label"] or "Entity"
        if d["accounts"]:
            for a in d["accounts"]:
                expected.append([drill["section"], "1", "ASSETS", layer, d["entity"],
                                 a["main_account"], a["account_name"], a["amount"]])
        else:
            expected.append([drill["section"], "1", "ASSETS", layer, d["entity"],
                             None, None, d["amount"]])
    assert body == expected
    total = _row_starting([r[3:] for r in rows if r[1] == "1"], "Total")
    assert total[4] == drill["total"]
    assert round(sum(r[7] for r in body), 2) == drill["total"]


def test_drill_sheet_follows_statement_heading_order():
    rows = _rows(_book()["Drill"])
    order = []
    for r in rows[2:]:
        if r[1] and r[1] not in order:
            order.append(r[1])
    assert order == ["4", "1", "2", "3"]


def test_a_statement_heading_with_no_drill_raises():
    drills = _drills()
    del drills["3"]
    with pytest.raises(ValueError, match="3 has no drill"):
        M.workbook(_payload(), drills)


def test_journals_sheet_lists_the_topside_journals():
    drill = _drills()["1"]
    topside = next(r for r in drill["rows"] if r["label"] == "Top-side journals")
    journal = topside["journals"][0]
    rows = _rows(_book()["Journals"])
    assert rows[1] == ["Heading", "Heading name", "Journal", "Description", "Amount",
                       "Posted by", "Approved by", "Basis"]
    assert rows[2] == ["1", "ASSETS", journal["journal_id"], journal["description"],
                       journal["amount"], journal["posted_by"], journal["approved_by"],
                       topside["journals_basis"]]


# --- drift: the texts shared with close-ui's numbers.js -----------------------

def _js():
    with open(NUMBERS_JS) as fh:
        return fh.read()


def test_signoff_labels_match_numbers_js():
    js = _js()
    block = js[js.index("const SIGNOFF_LABEL"):]
    block = block[:block.index("};")]
    found = dict(re.findall(r'(\w+):\s*\{\s*text:\s*"([^"]+)"', block))
    assert found == M.SIGNOFF_LABEL


def test_columns_match_numbers_js():
    js = _js()
    block = js[js.index("const COLUMNS"):]
    block = block[:block.index("};")]
    pl = re.search(r"\[PL\]:\s*(\[[^\]]*\])", block).group(1)
    bs = re.search(r"\[BS\]:\s*(\[[^\]]*\])", block).group(1)
    assert ast.literal_eval(pl) == M.COLUMNS["Profit and Loss"]
    assert ast.literal_eval(bs) == M.COLUMNS["Balance Sheet"]


def test_not_loaded_matches_numbers_js():
    assert '"%s"' % M.NOT_LOADED in _js()


def test_module_imports_no_frappe_and_no_konsol():
    with open(_PATH) as fh:
        tree = ast.parse(fh.read())
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            assert not any(a.name.startswith(("frappe", "konsol")) for a in node.names)
        if isinstance(node, ast.ImportFrom):
            assert not (node.module or "").startswith(("frappe", "konsol"))


def test_the_payload_is_not_mutated():
    payload = _payload()
    before = copy.deepcopy(payload)
    M.workbook(payload, _drills())
    assert payload == before


def test_account_lines_that_round_away_from_their_row_get_a_rounding_line():
    """Measured live 6 Oct (FY2025 P07 ECL_GROUP): each account amount is
    rounded on its own, so an entity row's account lines drifted up to
    0.14 from the row amount and the Drill sheet did not foot to the
    heading total. A rounding line per such row makes the column sum to
    the total, and says so."""
    drills = _drills()
    entity_row = drills["1"]["rows"][0]
    entity_row["accounts"] = [
        {"main_account": "1110", "account_name": "Cash", "amount": 150.01},
        {"main_account": "1120", "account_name": "Receivables", "amount": 150.0},
    ]  # row amount stays 300.0
    rows = _rows(_book(drills=drills)["Drill"])
    body = [r for r in rows[2:] if r[1] == "1" and r[3] != "Total"]
    rounding = [r for r in body if r[6] == "Rounding (accounts to row)"]
    assert len(rounding) == 1
    assert rounding[0][4] == entity_row["entity"]
    assert rounding[0][7] == -0.01
    assert round(sum(r[7] for r in body), 2) == drills["1"]["total"]
