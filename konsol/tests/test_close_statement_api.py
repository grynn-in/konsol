"""Statement read endpoint: konsol/close/statement_api.py ``get_statement``
(konsol#305 N51; stories 8.1, 8.3; #305-W4-1, W4-4, W4-5; W4-E8, W4-E10,
W4-E19).

Loaded against a stub frappe (pattern: test_close_ic_api.py's ``_Site`` /
``_frappe`` / ``_call``, copied, not imported). ``konsol.close.ch_read`` is
a stub module (no ClickHouse client involved); ``konsol.close.signoff_gate``
is a stub whose ``statement_accounts`` calls the REAL ``close_policy_model``
so the gap text is never hand-typed; ``konsol.close.statement_model`` is
the real pure module, loaded by path.
"""
import ast
import importlib.util
import json
import os
import re
import sys
import types
from datetime import date, datetime

import pytest

APP_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CLOSE_DIR = os.path.join(APP_DIR, "close")
API_PY = os.path.join(CLOSE_DIR, "statement_api.py")
FIXTURE_PY = os.path.join(APP_DIR, "tests", "fixtures", "close_statement_payload.json")
FIXTURE_DRILL_PY = os.path.join(APP_DIR, "tests", "fixtures", "close_drill_payload.json")
SITE_TZ = "Europe/London"


# --- N53: a static alias/column-collision check over EVERY ``*_SQL``
# constant in statement_api.py (konsol#305 N53; the family of bugs N52's
# live check found: an aggregate's own output alias also names a raw
# column referenced inside a SIBLING aggregate in the same SELECT list,
# which ClickHouse substitutes textually and then refuses as a nested
# aggregate — ``ILLEGAL_AGGREGATION`` or ``CYCLIC_ALIASES`` depending on
# the shape, per CLAUDE.md's ClickHouse traps section. This never needs a
# live ClickHouse or a stub frappe: it reads the SQL text straight from
# the source file with ``ast``, so it also guards every ``*_SQL`` constant
# added to this module later.)

def _sql_constants(path=API_PY):
    """``{name: sql}`` for every module-level ``..._SQL = (...)`` string
    constant, read with ``ast.literal_eval`` (handles the implicit
    string-literal concatenation each one uses) — no import required."""
    with open(path) as f:
        tree = ast.parse(f.read(), path)
    out = {}
    for node in tree.body:
        if not isinstance(node, ast.Assign) or len(node.targets) != 1:
            continue
        target = node.targets[0]
        if isinstance(target, ast.Name) and "_SQL" in target.id:
            out[target.id] = ast.literal_eval(node.value)
    return out


_IDENT_RE = re.compile(r"[A-Za-z_][A-Za-z0-9_]*")
_SQL_KEYWORDS = {"is", "null", "not", "and", "or", "distinct", "as"}
_AGG_COL_RE = re.compile(
    r"^([A-Za-z_][A-Za-z0-9_]*)\((.*)\)(?:\s+AS\s+([A-Za-z_][A-Za-z0-9_]*))?$",
    re.IGNORECASE | re.DOTALL)
_PLAIN_COL_RE = re.compile(
    r"^([A-Za-z_][A-Za-z0-9_]*)(?:\s+AS\s+([A-Za-z_][A-Za-z0-9_]*))?$", re.IGNORECASE)


def _split_select_list(sql):
    """The SELECT clause's top-level, comma-separated column expressions
    (commas inside a function call's parens don't split)."""
    upper = sql.upper()
    start = upper.index("SELECT") + len("SELECT")
    end = upper.index(" FROM ")
    parts, depth, current = [], 0, ""
    for ch in sql[start:end]:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        if ch == "," and depth == 0:
            parts.append(current)
            current = ""
        else:
            current += ch
    if current.strip():
        parts.append(current)
    return [p.strip() for p in parts]


def _column_entries(sql):
    """One ``(alias, is_aggregate, raw_identifiers)`` per SELECT column:
    ``raw_identifiers`` is what an aggregate call references by name
    (its own alias is never in this set), the thing ClickHouse will
    substitute another column's matching alias into."""
    entries = []
    for col in _split_select_list(sql):
        m = _AGG_COL_RE.match(col)
        if m:
            func, args, alias = m.group(1), m.group(2), m.group(3)
            idents = {t.lower() for t in _IDENT_RE.findall(args)}
            idents -= _SQL_KEYWORDS
            idents.discard(func.lower())
            entries.append((alias, True, idents))
            continue
        m = _PLAIN_COL_RE.match(col)
        if m:
            name, alias = m.group(1), m.group(2)
            entries.append((alias or name, False, set()))
        else:
            entries.append((None, False, set()))
    return entries


def _alias_collision(sql):
    """The first alias (or ``None``) that also names a raw column used
    inside a DIFFERENT aggregate's arguments in the same SELECT list —
    the collision ClickHouse raises ILLEGAL_AGGREGATION/CYCLIC_ALIASES
    over (measured live 4 Oct, N52/N53)."""
    entries = _column_entries(sql)
    for i, (alias, _is_agg, _idents) in enumerate(entries):
        if not alias:
            continue
        for j, (_alias2, is_agg2, idents2) in enumerate(entries):
            if j != i and is_agg2 and alias.lower() in idents2:
                return alias
    return None

LEAD = "zz-lead@example.com"
ANALYST = "zz-analyst@example.com"
VIEWER = "zz-viewer@example.com"
ENTITY_ACC = "zz-entity@example.com"


def _load_path(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


CPM = _load_path("close_policy_model_for_statement_api_test",
                  os.path.join(CLOSE_DIR, "close_policy_model.py"))
sys.modules.pop("close_policy_model_for_statement_api_test", None)


# --- the live-shaped chart (N41/N48/N49 facts; W41's codes adopted) --------
#
# Heading 1 ASSETS (Debit), 2 LIABILITIES (Credit), 3 EQUITY (Credit) on the
# Balance Sheet; heading 4 on the Profit and Loss. Leaves 2100 (under 2) and
# 4100 (under 4) are W41's fallback codes (konsol/tests/fixtures/
# close_approvals_journal_item.json), adopted here (W4-E19) so U41/W42 share
# one chart with the approvals screen's golden fixture.

def _accounts():
    return {
        "1": {"name": "1", "status": "Published", "account_name": "ASSETS", "parent_account": None, "is_group": 1,
              "statement_section": "Balance Sheet", "lft": 1, "normal_balance": "Debit"},
        "2": {"name": "2", "status": "Published", "account_name": "LIABILITIES", "parent_account": None, "is_group": 1,
              "statement_section": "Balance Sheet", "lft": 5, "normal_balance": "Credit"},
        "3": {"name": "3", "status": "Published", "account_name": "EQUITY", "parent_account": None, "is_group": 1,
              "statement_section": "Balance Sheet", "lft": 9, "normal_balance": "Credit"},
        "4": {"name": "4", "status": "Published", "account_name": "COST OF SALES", "parent_account": None, "is_group": 1,
              "statement_section": "Profit and Loss", "lft": 13, "normal_balance": ""},
        "1110": {"name": "1110", "status": "Published", "account_name": "Cash", "parent_account": "1", "is_group": 0,
                  "statement_section": "Balance Sheet", "lft": 2, "normal_balance": ""},
        "2100": {"name": "2100", "status": "Published", "account_name": "Accounts payable", "parent_account": "2",
                  "is_group": 0, "statement_section": "Balance Sheet", "lft": 6,
                  "normal_balance": ""},
        "3200": {"name": "3200", "status": "Published", "account_name": "Share capital", "parent_account": "3",
                  "is_group": 0, "statement_section": "Balance Sheet", "lft": 10,
                  "normal_balance": ""},
        "3300": {"name": "3300", "status": "Published", "account_name": "AOCI - CTA", "parent_account": "3",
                  "is_group": 0, "statement_section": "Balance Sheet", "lft": 11,
                  "normal_balance": ""},
        "3100": {"name": "3100", "status": "Published", "account_name": "Retained earnings", "parent_account": "3",
                  "is_group": 0, "statement_section": "Balance Sheet", "lft": 12,
                  "normal_balance": ""},
        "4100": {"name": "4100", "status": "Published", "account_name": "Sales revenue", "parent_account": "4",
                  "is_group": 0, "statement_section": "Profit and Loss", "lft": 14,
                  "normal_balance": ""},
    }


_DECLARED_ROWS = {
    "3300": {"is_group": 0, "status": "Published", "statement_section": "Balance Sheet",
              "account_name": "AOCI - CTA", "parent_account": "3"},
    "3100": {"is_group": 0, "status": "Published", "statement_section": "Balance Sheet",
              "account_name": "Retained earnings", "parent_account": "3"},
}


def _calendar(years):
    rows = []
    for fy in years:
        rows.append(_period_row(fy, 0, "Opening", "Closed"))
        for fp in range(1, 13):
            rows.append(_period_row(fy, fp, "Regular", "Closed"))
        rows.append(_period_row(fy, 13, "Closing", "Closed"))
    return rows


def _period_row(fy, fp, period_type, status):
    month = min(max(fp, 1), 12)
    return {
        "fiscal_year": fy, "fiscal_period": fp,
        "period_code": "FY%dP%02d" % (fy, fp), "period_label": "FY%d P%02d" % (fy, fp),
        "period_type": period_type, "start_date": date(fy, month, 1),
        "end_date": date(fy, month, 28), "quarter": "Q%d" % ((month - 1) // 3 + 1),
        "status": status,
    }


_CALENDAR = _calendar((2024, 2025))
_KEY = (2025, 7)


def _tb_row(main_account, amount, adjustment_type=None, fy=2025, fp=7, null_rows=0):
    return {"fiscal_year": fy, "fiscal_period": fp, "main_account": main_account,
            "adjustment_type": adjustment_type, "amount": amount, "null_rows": null_rows}


def _drill_tb_row(main_account, amount, adjustment_type, data_area_id, fy=2025, fp=7, null_rows=0):
    """N52's entity-grain read: like ``_tb_row`` but carrying
    ``data_area_id``, and keyed ``amt`` (not ``amount``) — the real SQL's
    own alias, chosen to avoid ClickHouse's alias/column collision with
    ``countIf(amount IS NULL)`` (measured live 4 Oct, ILLEGAL_AGGREGATION);
    the product's ``_drill_row`` renames it back before use."""
    return {"fiscal_year": fy, "fiscal_period": fp, "data_area_id": data_area_id,
            "main_account": main_account, "adjustment_type": adjustment_type,
            "amt": amount, "null_rows": null_rows}


def _drill_journal_row(journal_id, main_account, amount, data_area_id, description,
                        posted_by, approved_by, adjustment_type="topside"):
    """N52's journals read: the ClickHouse column is aliased ``amount`` (the
    SQL sums ``net_amount``); the product maps it to drill_model's
    ``net_amount`` key before calling ``drill_model.drill``."""
    return {"journal_id": journal_id, "adjustment_type": adjustment_type,
            "data_area_id": data_area_id, "main_account": main_account,
            "amount": amount, "description": description,
            "posted_by": posted_by, "approved_by": approved_by}


class _Site:
    def __init__(self):
        self.user = ANALYST
        self.roles = {"EPM Analyst"}
        self.periods = [dict(r) for r in _CALENDAR]
        self._open(2025, 7)
        self.groups = [{"consolidation_group": "G1", "reporting_currency": "USD"}]
        self.accounts = _accounts()
        self.cta_account, self.result_account = "3300", "3100"
        self.declared_rows = dict(_DECLARED_ROWS)
        self.tb_rows = [
            _tb_row("1110", 1735.10),
            _tb_row("2100", -803.70),
            _tb_row("3200", -684.90),
            _tb_row("CTA", -5.07, adjustment_type="cta"),
            _tb_row("4100", -241.43),
        ]
        self.entity_rows = [{"data_area_id": "ZZA"}]
        self.in_scope = {"ZZA", "ZZB"}
        self.allowed = None
        self.ch_error = None
        self.ch_calls = []
        self.drill_tb_rows = []  # N52: get_drill's entity-grain read
        self.journal_rows = []  # N52: get_drill's gold_consolidation_adjustments read
        self.commentary_rows = [
            {"name": "SC-G1-2025-7-4", "consolidation_group": "G1", "fiscal_year": 2025,
             "fiscal_period": 7, "heading": "4", "text": "Strong quarter.",
             "modified": datetime(2025, 8, 1, 10, 0, 0), "modified_by": ANALYST},
        ]
        self.users = [{"name": ANALYST, "full_name": "Zz Analyst"}]
        self.run = None  # latest_close_run: None -> "provisional"
        self.reads = []

    def _open(self, fy, fp):
        for row in self.periods:
            if (row["fiscal_year"], row["fiscal_period"]) == (fy, fp):
                row["status"] = "Open"


def _match_value(value, cond):
    if isinstance(cond, (list, tuple)):
        op, arg = cond[0], cond[1]
        if op == "in":
            return value in arg
        if op == "is":
            assert arg in ("set", "not set"), cond
            return (value not in (None, "")) == (arg == "set")
        raise AssertionError("stub: unsupported operator %r" % (op,))
    return value == cond


def _match(row, filters):
    return all(_match_value(row.get(key), cond) for key, cond in (filters or {}).items())


def _frappe(site):
    frappe = types.ModuleType("frappe")
    frappe.ValidationError = type("ValidationError", (Exception,), {})
    frappe.PermissionError = type("PermissionError", (Exception,), {})

    def throw(msg, exc=None, **k):
        raise (exc or frappe.ValidationError)(msg)

    def only_for(roles, message=False):
        roles = [roles] if isinstance(roles, str) else list(roles)
        if not site.roles.intersection(roles):
            raise frappe.PermissionError("Not permitted")

    def whitelist(*a, **k):
        return lambda fn: fn

    def get_all(doctype, filters=None, fields=None, limit_page_length=None, **k):
        site.reads.append(("get_all", doctype))
        if doctype == "Consolidation Group":
            rows = [r for r in site.groups if _match(r, filters)]
        elif doctype == "Main Account":
            rows = [r for r in site.accounts.values() if _match(r, filters)]
        elif doctype == "Statement Commentary":
            rows = [r for r in site.commentary_rows if _match(r, filters)]
        elif doctype == "User":
            rows = [r for r in site.users if _match(r, filters)]
        else:
            raise AssertionError("unexpected get_all on %s" % doctype)
        return [{f: r.get(f) for f in fields} for r in rows]

    frappe.throw = throw
    frappe.only_for = only_for
    frappe.whitelist = whitelist
    frappe.get_all = get_all
    frappe.get_roles = lambda user=None: sorted(site.roles)
    frappe.session = types.SimpleNamespace(user=site.user)
    site.response = {}
    frappe.response = site.response
    frappe.utils = types.SimpleNamespace(get_system_timezone=lambda: SITE_TZ)
    return frappe


def _ch_read(site):
    import re

    ch = types.ModuleType("konsol.close.ch_read")
    _ERROR_NAME = re.compile(r"\(([A-Z][A-Z0-9_]+)\)")
    _NOT_BUILT = {"UNKNOWN_TABLE", "UNKNOWN_DATABASE"}

    def rows(sql, params=None):
        site.ch_calls.append((sql, dict(params or {})))
        if site.ch_error is not None:
            raise site.ch_error
        if "gold_consolidation_adjustments" in sql:
            return [dict(r) for r in site.journal_rows]
        if "DISTINCT data_area_id" in sql:
            return [dict(r) for r in site.entity_rows]
        if "data_area_id" in sql:  # N52: the entity-grain drill read
            return [dict(r) for r in site.drill_tb_rows]
        return [dict(r) for r in site.tb_rows]

    ch.rows = rows
    ch.error_names = lambda e: set(_ERROR_NAME.findall(str(e)))
    ch.not_built = lambda e: bool(ch.error_names(e) & _NOT_BUILT)
    return ch


def _signoff_gate(site):
    sg = types.ModuleType("konsol.close.signoff_gate")
    sg.statement_accounts = lambda: CPM.statement_accounts(
        site.cta_account, site.result_account, site.declared_rows)
    sg.in_scope_entities = lambda fy, fp: sorted(site.in_scope)
    return sg


def _assertion_run(site):
    mod = types.ModuleType("konsol.consolidation.doctype.assertion_run.assertion_run")
    mod.SIGNED_STATES = ("Signed Off", "Acknowledged", "Overridden")
    mod.RE_SIGN_NEEDED = "Re-sign Needed"

    def latest_close_run(fy, fp):
        site.reads.append(("run", fy, fp))
        return dict(site.run) if site.run else None

    mod.latest_close_run = latest_close_run
    return mod


def _invoke(site, run):
    frappe = _frappe(site)
    konsol = types.ModuleType("konsol")
    konsol.__path__ = []
    close = types.ModuleType("konsol.close")
    close.__path__ = []
    fiscal_calendar = types.ModuleType("konsol.fiscal_calendar")
    fiscal_calendar.fiscal_period_rows = lambda: [dict(r) for r in site.periods]
    entity_permissions = types.ModuleType("konsol.entity_permissions")
    entity_permissions.allowed_entity_codes = lambda user=None: site.allowed
    consolidation = types.ModuleType("konsol.consolidation")
    doctype_pkg = types.ModuleType("konsol.consolidation.doctype")
    assertion_run_pkg = types.ModuleType("konsol.consolidation.doctype.assertion_run")
    assertion_run_mod = _assertion_run(site)

    konsol.close = close
    konsol.fiscal_calendar = fiscal_calendar
    konsol.entity_permissions = entity_permissions
    konsol.consolidation = consolidation
    close.ch_read = _ch_read(site)
    close.signoff_gate = _signoff_gate(site)

    names = ["frappe", "konsol", "konsol.close", "konsol.fiscal_calendar",
             "konsol.entity_permissions", "konsol.close.ch_read", "konsol.close.signoff_gate",
             "konsol.close.statement_model", "konsol.close.timefmt", "konsol.close.drill_model",
             "konsol.close.statement_export_model",
             "konsol.consolidation", "konsol.consolidation.doctype",
             "konsol.consolidation.doctype.assertion_run",
             "konsol.consolidation.doctype.assertion_run.assertion_run",
             "close_statement_api_under_test"]
    saved = {n: sys.modules.get(n) for n in names}
    sys.modules.update({
        "frappe": frappe, "konsol": konsol, "konsol.close": close,
        "konsol.fiscal_calendar": fiscal_calendar,
        "konsol.entity_permissions": entity_permissions,
        "konsol.close.ch_read": close.ch_read, "konsol.close.signoff_gate": close.signoff_gate,
        "konsol.consolidation": consolidation, "konsol.consolidation.doctype": doctype_pkg,
        "konsol.consolidation.doctype.assertion_run": assertion_run_pkg,
        "konsol.consolidation.doctype.assertion_run.assertion_run": assertion_run_mod,
    })
    try:
        close.statement_model = _load_path(
            "konsol.close.statement_model", os.path.join(CLOSE_DIR, "statement_model.py"))
        close.timefmt = _load_path(
            "konsol.close.timefmt", os.path.join(CLOSE_DIR, "timefmt.py"))
        close.drill_model = _load_path(
            "konsol.close.drill_model", os.path.join(CLOSE_DIR, "drill_model.py"))
        close.statement_export_model = _load_path(
            "konsol.close.statement_export_model",
            os.path.join(CLOSE_DIR, "statement_export_model.py"))
        api = _load_path("close_statement_api_under_test", API_PY)
        return run(api)
    finally:
        for n, old in saved.items():
            if old is None:
                sys.modules.pop(n, None)
            else:
                sys.modules[n] = old


def _call(site, fy=2025, fp=7, group=None):
    result = _invoke(site, lambda api: api.get_statement(fy, fp, group))
    json.dumps(result)  # JSON-safe
    return result


def _call_raises(site, fy=2025, fp=7, group=None):
    with pytest.raises(Exception) as info:
        _call(site, fy, fp, group)
    return info.value


def _call_drill(site, fy=2025, fp=7, group="G1", heading="1"):
    result = _invoke(site, lambda api: api.get_drill(fy, fp, group, heading))
    json.dumps(result)  # JSON-safe
    return result


def _call_drill_raises(site, fy=2025, fp=7, group="G1", heading="1"):
    with pytest.raises(Exception) as info:
        _call_drill(site, fy, fp, group, heading)
    return info.value


def _bs_line(result, heading):
    for line in result["statement"]["sections"][1]["lines"]:
        if line.get("heading") == heading:
            return line
    raise AssertionError("no BS line for %r" % heading)


# --- happy path -------------------------------------------------------------

def test_happy_path_returns_ok_with_both_sections():
    site = _Site()
    result = _call(site)
    assert result["state"] == "ok"
    assert [s["section"] for s in result["statement"]["sections"]] == [
        "Profit and Loss", "Balance Sheet"]
    residual = result["statement"]["sections"][1]["lines"][-1]
    assert residual["kind"] == "residual"
    assert residual["current"] == 0.0
    assert result["signoff"] == {"state": "provisional", "run": None}
    assert result["can_comment"] is True
    assert result["commentary"]["4"]["text"] == "Strong quarter."
    assert result["not_included"]["count"] == 1
    assert result["consolidation_group"] == "G1"
    assert result["group_note"] == "The only consolidation group."
    assert result["gap"] is None


def test_query_count_is_independent_of_heading_and_entity_count():
    site = _Site()
    _call(site)
    first = len(site.reads)
    site2 = _Site()
    for n in range(3, 31):
        code = "99%02d" % n
        site2.accounts[code] = {"name": code, "status": "Published", "account_name": "X%d" % n,
                                "parent_account": "1", "is_group": 0,
                                "statement_section": "Balance Sheet", "lft": 100 + n,
                                "normal_balance": ""}
        site2.in_scope.add("ZZ%02d" % n)
    _call(site2)
    assert len(site2.reads) == first
    assert len(site2.ch_calls) == len(site.ch_calls) == 2


def test_no_sql_names_a_dim_column():
    site = _Site()
    _call(site)
    for sql, _params in site.ch_calls:
        assert "dim_" not in sql


def test_no_sql_constant_has_an_alias_column_collision():
    """N53: every ``*_SQL`` constant in statement_api.py, read straight
    from source — no alias may also name a raw column referenced inside
    a sibling aggregate in the same SELECT list (the ILLEGAL_AGGREGATION/
    CYCLIC_ALIASES family). Stub-based tests can never catch this:
    ``ch_read`` is stubbed everywhere else in this file, so it never runs
    real SQL against real ClickHouse."""
    constants = _sql_constants()
    assert set(constants) == {
        "_TB_SQL", "_ENTITY_SQL", "_DRILL_SQL", "_DRILL_SQL_WITH_CTA",
        "_DRILL_JOURNALS_SQL",
    }, "a *_SQL constant was added or removed: update this test's expected set"
    for name, sql in constants.items():
        collision = _alias_collision(sql)
        assert collision is None, (
            "%s: alias %r collides with a raw column referenced inside a "
            "sibling aggregate in the same SELECT list (measured live as "
            "ILLEGAL_AGGREGATION/CYCLIC_ALIASES) — alias the aggregate "
            "differently (e.g. AS amt) and remap it in Python before use."
            % (name, collision)
        )


# --- choose a group / unknown group (W4-E8) --------------------------------

def test_two_groups_and_no_group_named_is_choose_group_with_no_clickhouse_read():
    site = _Site()
    site.groups = [{"consolidation_group": "G1", "reporting_currency": "USD"},
                   {"consolidation_group": "G2", "reporting_currency": "EUR"}]
    result = _call(site)
    assert result["state"] == "choose_group"
    assert result["consolidation_group"] is None
    assert result["statement"] is None
    assert site.ch_calls == []


def test_unknown_group_throws():
    site = _Site()
    _call_raises(site, group="ZZ-NOPE")


# --- chart missing ----------------------------------------------------------

def test_empty_chart_is_no_chart_with_no_clickhouse_read():
    site = _Site()
    site.accounts = {}
    result = _call(site)
    assert result["state"] == "no_chart"
    assert result["statement"] is None
    assert site.ch_calls == []


# --- warehouse failures (never an empty statement) --------------------------

def test_warehouse_down_is_an_error_state_naming_the_exception():
    site = _Site()
    site.ch_error = RuntimeError("boom")
    result = _call(site)
    assert result["state"] == "error"
    assert "RuntimeError" in result["message"]
    assert result["statement"] is None


def test_unbuilt_relation_is_not_built():
    site = _Site()
    site.ch_error = RuntimeError("Code: 60. DB::Exception: Table (UNKNOWN_TABLE)")
    result = _call(site)
    assert result["state"] == "not_built"
    assert result["statement"] is None


def test_null_amount_is_an_error_state_naming_the_account():
    site = _Site()
    site.tb_rows.append(_tb_row("1110", None, null_rows=1))
    result = _call(site)
    assert result["state"] == "error"
    assert "1110" in result["message"]
    assert result["statement"] is None


def test_bs_heading_with_no_normal_balance_is_a_visible_setup_gap_never_a_500():
    """S5: a BS heading with a blank ``normal_balance`` never crashes the
    endpoint and never silently flips a side — it comes back as a named,
    visible SETUP gap, never the same state as a warehouse failure (a
    Close Lead's job, not an outage)."""
    site = _Site()
    site.accounts["2"]["normal_balance"] = ""
    result = _call(site)
    assert result["state"] == "setup_gap"
    assert "statement_heading_side_undeclared" in result["message"]
    assert "2" in result["message"]
    assert result["statement"] is None


# --- golden payload (S5, W4-E19): the real producer's setup_gap output ------

FIXTURE_SETUP_GAP_PY = os.path.join(
    APP_DIR, "tests", "fixtures", "close_statement_payload_setup_gap.json")


def test_golden_setup_gap_payload_matches_the_committed_fixture():
    site = _Site()
    site.accounts["2"]["normal_balance"] = ""
    result = json.loads(json.dumps(_call(site)))  # the wire shape: no tuples
    with open(FIXTURE_SETUP_GAP_PY) as fh:
        expected = json.load(fh)
    assert result == expected


# --- undeclared statement accounts: never silently balanced ----------------

def test_undeclared_statement_accounts_gap_present_residual_non_zero():
    site = _Site()
    site.cta_account, site.result_account = "", ""
    result = _call(site)
    assert result["state"] == "ok"
    assert result["gap"]["code"] == "statement_accounts_undeclared"
    residual = result["statement"]["sections"][1]["lines"][-1]
    assert residual["current"] != 0.0
    assert "explained" in residual


# --- can_comment (#305-W4-5 5b) ---------------------------------------------

def test_viewer_cannot_comment():
    site = _Site()
    site.user, site.roles = VIEWER, {"EPM User"}
    result = _call(site)
    assert result["can_comment"] is False


def test_closed_period_cannot_comment():
    site = _Site()
    for row in site.periods:
        if (row["fiscal_year"], row["fiscal_period"]) == _KEY:
            row["status"] = "Closed"
    result = _call(site)
    assert result["can_comment"] is False


def test_entity_accountant_is_refused_by_only_for():
    site = _Site()
    site.user, site.roles = ENTITY_ACC, {"Entity Accountant"}
    _call_raises(site)


def test_an_undeclared_period_throws():
    site = _Site()
    _call_raises(site, fy=2099, fp=1)


# --- sign-off label follows the run (#305-W4-4) -----------------------------

def test_signed_run_gives_signed_state():
    site = _Site()
    site.run = {"name": "AR-1", "signoff_status": "Signed Off"}
    result = _call(site)
    assert result["signoff"] == {"state": "signed", "run": "AR-1"}


def test_re_sign_needed_run_gives_resign_needed_state():
    site = _Site()
    site.run = {"name": "AR-2", "signoff_status": "Re-sign Needed"}
    result = _call(site)
    assert result["signoff"] == {"state": "resign_needed", "run": "AR-2"}


# --- entity scope (W4-E9/E10) ------------------------------------------------

def test_scoped_caller_sees_no_names_only_a_hidden_count():
    site = _Site()
    site.allowed = ["ZZA"]  # ZZB is missing and out of scope: hidden, never named
    result = _call(site)
    assert result["not_included"]["entities"] == []
    assert result["not_included"]["hidden"] == 1
    assert result["not_included"]["count"] == 1


# --- golden payload (W4-E19): the real producer's output, never hand-built --

def test_golden_payload_matches_the_committed_fixture():
    site = _Site()
    result = json.loads(json.dumps(_call(site)))  # the wire shape: no tuples
    with open(FIXTURE_PY) as fh:
        expected = json.load(fh)
    assert result == expected


# =============================================================================
# get_drill (konsol#305 N52): one heading's amount, entity grain, with
# journals for the top-side row. Same stub site, extended with the
# entity-grain read (``drill_tb_rows``) and the journals read
# (``journal_rows``).
# =============================================================================

def _drill_site():
    """Heading "1" (ASSETS, Debit) carries the CTA account for this test
    (a leaf "1200" added under it), so one drill shows entity, IC, CTA and
    top-side rows together, as the row's test-first bullet asks (N50's own
    layer coverage, combined onto one heading here). The aggregate
    (``tb_rows``) and entity-grain (``drill_tb_rows``) totals agree by
    construction: 300 (ZZA) + 200 (ZZB) + 100 (ZZC) - 50 (IC) + 40
    (topside) = 590 on 1110; CTA -5 + -3 = -8."""
    site = _Site()
    site.accounts["1200"] = {
        "name": "1200", "status": "Published", "account_name": "CTA reserve",
        "parent_account": "1", "is_group": 0, "statement_section": "Balance Sheet",
        "lft": 3, "normal_balance": "",
    }
    site.cta_account = "1200"
    site.declared_rows["1200"] = {"is_group": 0, "status": "Published",
                                   "statement_section": "Balance Sheet",
                                   "account_name": "CTA reserve",
                                   "parent_account": "1"}
    site.tb_rows = [
        _tb_row("1110", 590.0),
        _tb_row("CTA", -8.0, adjustment_type="cta"),
        _tb_row("2100", -803.70),
        _tb_row("3200", -684.90),
        _tb_row("4100", -241.43),
    ]
    site.drill_tb_rows = [
        _drill_tb_row("1110", 300.0, "entity", "ZZA"),
        _drill_tb_row("1110", 200.0, "entity", "ZZB"),
        _drill_tb_row("1110", 100.0, "entity", "ZZC"),
        _drill_tb_row("1110", -50.0, "ic_elimination", ""),
        _drill_tb_row("1110", 40.0, "topside", "ZZA"),
        # #305 story 6.5: an earlier period's journal reversing here, and a
        # topside of the same size, so every existing total stays as it was.
        _drill_tb_row("1110", 10.0, "topside", "ZZB"),
        _drill_tb_row("1110", -10.0, "auto_reversal", "ZZA"),
        _drill_tb_row("CTA", -5.0, "cta", "ZZA"),
        _drill_tb_row("CTA", -3.0, "cta", "ZZB"),
    ]
    site.journal_rows = [
        _drill_journal_row("J-1", "1110", 40.0, "ZZA", "Reclass intercompany loan",
                            "alice@example.com", "bob@example.com"),
        _drill_journal_row("J-2", "1110", 10.0, "ZZB", "Accrue audit fee",
                            "alice@example.com", "bob@example.com"),
        _drill_journal_row("J-P6", "1110", -10.0, "ZZA", "June bonus",
                            "alice@example.com", "bob@example.com",
                            adjustment_type="auto_reversal"),
    ]
    return site


def test_get_drill_returns_entity_ic_cta_and_topside_rows_for_heading_1():
    site = _drill_site()
    result = _call_drill(site)
    assert result["state"] == "ok"
    assert result["heading"] == "1"
    assert result["consolidation_group"] == "G1"
    drill = result["drill"]
    assert drill["heading"] == "1"
    assert drill["total"] == 582.0  # 590 (leaf, incl. IC/topside) - 8 (CTA)

    entity_rows = [r for r in drill["rows"] if r["layer"] == "entity" and r["entity"]]
    assert {r["entity"]: r["amount"] for r in entity_rows} == {
        "ZZA": 300.0, "ZZB": 200.0, "ZZC": 100.0}

    ic_rows = [r for r in drill["rows"] if r["label"] == "Intercompany eliminations"]
    assert len(ic_rows) == 1
    assert ic_rows[0]["amount"] == -50.0
    assert ic_rows[0]["entity"] is None

    topside_rows = [r for r in drill["rows"] if r["label"] == "Top-side journals"]
    assert len(topside_rows) == 1
    assert topside_rows[0]["amount"] == 40.0
    assert topside_rows[0]["source"] == {"kind": "journals"}
    assert topside_rows[0]["journals"][0]["journal_id"] == "J-1"
    assert topside_rows[0]["journals"][0]["amount"] == 40.0

    cta_rows = [r for r in drill["rows"] if r["layer"] == "cta"]
    assert len(cta_rows) == 1
    assert cta_rows[0]["amount"] == -8.0
    assert cta_rows[0]["entity"] is None


# --- failure path: warehouse down -------------------------------------------

def test_warehouse_error_in_drill_is_an_error_state_with_no_drill():
    site = _drill_site()
    site.ch_error = RuntimeError("boom")
    result = _call_drill(site)
    assert result["state"] == "error"
    assert "RuntimeError" in result["message"]
    assert result["drill"] is None


# --- failure path: heading must be a Published group heading ----------------

def test_a_leaf_code_as_heading_throws():
    site = _drill_site()
    _call_drill_raises(site, heading="1110")


def test_an_unknown_code_as_heading_throws():
    site = _drill_site()
    _call_drill_raises(site, heading="ZZ-NOPE")


def test_an_unknown_group_throws_for_get_drill():
    site = _drill_site()
    _call_drill_raises(site, group="ZZ-NOPE")


# --- no dim_ column; exactly 3 ClickHouse calls ------------------------------

def test_no_drill_sql_names_a_dim_column_and_exactly_three_calls_are_made():
    site = _drill_site()
    _call_drill(site)
    assert len(site.ch_calls) == 3
    for sql, _params in site.ch_calls:
        assert "dim_" not in sql


# --- scope (W4-E9): hidden entities are never named --------------------------

def test_scoped_caller_sees_the_aggregated_outside_scope_row():
    site = _drill_site()
    site.allowed = ["ZZA"]
    result = _call_drill(site)
    rows = result["drill"]["rows"]
    outside = next(r for r in rows if r["layer"] == "entity" and r["entity"] is None)
    assert outside["label"] == "2 entities outside your scope"
    assert outside["amount"] == 300.0  # 200 (ZZB) + 100 (ZZC)
    assert outside["accounts"] == []
    for row in rows:
        assert row.get("entity") not in ("ZZB", "ZZC")


def test_the_drill_tags_the_auto_reversal_journal_row():
    """#305 story 6.5: the auto-reversal row keeps its original journal_id
    (V01), so the drill tags it."""
    result = _call_drill(_drill_site())
    topside = next(r for r in result["drill"]["rows"] if r["label"] == "Top-side journals")
    assert [(j["journal_id"], j["auto_reversal"]) for j in topside["journals"]] == \
        [("J-1", False), ("J-2", False), ("J-P6", True)]


# --- golden payload (W4-E19): the real producer's output --------------------

def test_golden_drill_payload_matches_the_committed_fixture():
    site = _drill_site()
    result = json.loads(json.dumps(_call_drill(site)))  # the wire shape: no tuples
    with open(FIXTURE_DRILL_PY) as fh:
        expected = json.load(fh)
    assert result == expected


# =============================================================================
# export_statement (konsol#305 story 8.5, decision #305-W5-3): the Numbers
# statement as a downloaded .xlsx — same roles and scoping as get_statement,
# a non-ok state refused with the server's own sentence, never an empty
# file. The workbook is opened back and checked against the REAL golden
# get_statement payload (close_statement_payload.json).
# =============================================================================

def _export_site():
    """``_Site`` (whose get_statement output IS the golden fixture), plus an
    entity-grain drill read that agrees with its aggregate TB rows, and one
    top-side journal: 1110 = 1695.10 (ZZA) + 40 (topside) = 1735.10."""
    site = _Site()
    site.drill_tb_rows = [
        _drill_tb_row("1110", 1695.10, "entity", "ZZA"),
        _drill_tb_row("1110", 40.0, "topside", "ZZA"),
        _drill_tb_row("2100", -803.70, "entity", "ZZA"),
        _drill_tb_row("3200", -684.90, "entity", "ZZA"),
        _drill_tb_row("CTA", -5.07, "cta", "ZZA"),
        _drill_tb_row("4100", -241.43, "entity", "ZZA"),
    ]
    site.journal_rows = [
        _drill_journal_row("J-1", "1110", 40.0, "ZZA", "Reclass intercompany loan",
                            "alice@example.com", "bob@example.com"),
    ]
    return site


def _call_export(site, fy=2025, fp=7, group=None):
    _invoke(site, lambda api: api.export_statement(fy, fp, group))
    return site.response


def _export_refusal(site, group=None, kind="ValidationError"):
    with pytest.raises(Exception) as info:
        _call_export(site, group=group)
    assert type(info.value).__name__ == kind, repr(info.value)
    assert "filecontent" not in site.response  # never an empty file
    return info.value


def _export_book(site, group=None):
    import io
    from openpyxl import load_workbook

    response = _call_export(site, group=group)
    return load_workbook(io.BytesIO(response["filecontent"]))


def _golden():
    with open(FIXTURE_PY) as fh:
        return json.load(fh)


def test_export_site_statement_is_the_golden_payload():
    site = _export_site()
    assert json.loads(json.dumps(_call(site))) == _golden()


def test_export_is_a_binary_xlsx_named_for_the_period_and_group():
    site = _export_site()
    response = _call_export(site)
    assert response["type"] == "binary"
    assert response["filename"] == "numbers-FY2025P07-G1.xlsx"
    assert response["filecontent"][:2] == b"PK"  # a zip container


def test_export_statement_sheet_matches_the_golden_payload():
    golden = _golden()
    book = _export_book(_export_site())
    rows = [list(r) for r in book["Statement"].iter_rows(values_only=True)]
    assert rows[0][0] == "Numbers · FY2025P07 · G1 · USD"
    assert rows[1][0] == "Provisional"
    assert rows[2][0] == golden["statement"]["legend"]
    by_code = {r[1]: r for r in rows if r[1]}
    for section in golden["statement"]["sections"]:
        for line in section["lines"]:
            if line["kind"] == "heading":
                assert by_code[line["heading"]][0] == line["heading_name"]
                assert by_code[line["heading"]][2] == line["current"]
    assert by_code["4"][6] == golden["commentary"]["4"]["text"]
    assert by_code["4"][7].startswith(golden["commentary"]["4"]["by"] + " · ")


def test_export_drill_sheet_totals_equal_every_golden_heading():
    golden = _golden()
    book = _export_book(_export_site())
    rows = [list(r) for r in book["Drill"].iter_rows(values_only=True)]
    totals = {r[1]: r[7] for r in rows if r[3] == "Total"}
    expected = {line["heading"]: line["current"]
                for section in golden["statement"]["sections"]
                for line in section["lines"] if line["kind"] == "heading"}
    assert totals == expected
    journals = [list(r) for r in book["Journals"].iter_rows(values_only=True)]
    assert journals[2][:5] == ["1", "ASSETS", "J-1", "Reclass intercompany loan", 40.0]


def test_export_reads_clickhouse_four_times_whatever_the_heading_count():
    site = _export_site()
    _call_export(site)
    assert len(site.ch_calls) == 4
    for sql, _params in site.ch_calls:
        assert "dim_" not in sql
    site2 = _export_site()
    for n in range(5, 9):
        code = str(n)
        site2.accounts[code] = {"name": code, "status": "Published", "account_name": "H%d" % n,
                                "parent_account": None, "is_group": 1,
                                "statement_section": "Balance Sheet", "lft": 100 + n,
                                "normal_balance": "Debit"}
    _call_export(site2)
    assert len(site2.ch_calls) == 4


def test_export_choose_group_is_refused_with_the_servers_sentence():
    site = _export_site()
    site.groups.append({"consolidation_group": "G2", "reporting_currency": "EUR"})
    exc = _export_refusal(site)
    assert str(exc) == "Choose a consolidation group."


def test_export_setup_gap_is_refused_with_the_servers_sentence():
    site = _export_site()
    site.accounts["2"]["normal_balance"] = ""
    expected = _call(site)["message"]
    exc = _export_refusal(site)
    assert str(exc) == expected
    assert expected.startswith("statement_heading_side_undeclared")


def test_export_not_built_and_error_are_refused_with_the_servers_sentence():
    for error in (RuntimeError("Code: 60 (UNKNOWN_TABLE)"), RuntimeError("boom")):
        site = _export_site()
        site.ch_error = error
        expected = _call(site)["message"]
        exc = _export_refusal(site)
        assert str(exc) == expected


def test_export_no_chart_is_refused_with_the_servers_sentence():
    site = _export_site()
    site.accounts = {}
    exc = _export_refusal(site)
    assert str(exc) == "Publish the group chart (Main Account) first."


def test_export_a_drill_that_disagrees_with_the_statement_is_refused():
    site = _export_site()
    site.drill_tb_rows[0] = _drill_tb_row("1110", 1.0, "entity", "ZZA")
    exc = _export_refusal(site)
    assert "does not match the statement" in str(exc)


def test_export_entity_accountant_is_refused_by_only_for():
    site = _export_site()
    site.roles = {"EPM Entity Accountant"}
    _export_refusal(site, kind="PermissionError")
    assert site.ch_calls == []


def test_export_scoped_caller_never_sees_hidden_entity_names():
    site = _export_site()
    site.drill_tb_rows.append(_drill_tb_row("1110", 0.0, "entity", "ZZB"))
    site.allowed = ["ZZA"]
    book = _export_book(site)
    for sheet in book.worksheets:
        for row in sheet.iter_rows(values_only=True):
            assert not any(isinstance(v, str) and "ZZB" in v for v in row), (sheet.title, row)
    layers = {r[3] for r in book["Drill"].iter_rows(values_only=True)}
    assert "1 entities outside your scope" in layers


def test_export_filename_names_the_fiscal_year_even_when_the_period_code_does_not():
    """Measured live 6 Oct: the live period_code is "P07", so a name built
    from it ("numbers-P07-ECL_GROUP.xlsx") lost the fiscal year."""
    site = _export_site()
    for row in site.periods:
        row["period_code"] = "P%02d" % row["fiscal_period"]
    response = _call_export(site, group="G1")
    assert response["filename"] == "numbers-FY2025P07-G1.xlsx"
