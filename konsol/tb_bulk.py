"""Bulk trial balance upload: one file, many entity-periods.

Flow (konsol-exec's upload page, or the Trial Balance Upload form):

  1. The file is uploaded as a private File, then `check_file` creates a
     Trial Balance Upload and reports on every entity-period in it. Nothing
     is loaded; the report says which are ready and why the others are not.
  2. `load` re-checks (a period may have closed, or someone submitted in
     between). If new problems appeared and the user did not choose to skip
     them, it returns the fresh report with a `refused` reason instead of
     loading. Otherwise it queues `run_load`.
  3. `run_load` (a background job, run as the user who asked) creates one
     ordinary Trial Balance Submission per ready entity-period and submits
     it. Each one commits on its own: its rows are in the warehouse once
     submitted, so a later failure must not roll it back. Progress is
     written to the upload as it goes.
  4. A load that stopped (worker restart, timeout) or finished partly can be
     resumed with `load` again: rows already loaded are carried forward and
     never loaded twice.

Only the Close Lead (EPM Admin) or a System Manager may load. An upload is
visible to a user only if they may see every entity in it (or uploaded it).
"""
import json
from io import BytesIO

import frappe
from frappe.utils import cint, strip_html

from konsol import period_status
from konsol import tb_bulk_model as M
from konsol.clickhouse import execute
from konsol.consolidation.doctype.trial_balance_submission.trial_balance_submission import (
    CONTROL_TABLE,
    _sql_str,
    validate_tb_rows,
)
from konsol.entity_permissions import allowed_entity_codes

DOCTYPE = "Trial Balance Upload"
TERMINAL = ("Loaded", "Partly Loaded", "Failed")
JOB_PREFIX = "konsol-tb-upload::"


def _job_id(name):
    return f"{JOB_PREFIX}{name}"


def _job_running(name):
    """True while the load job is queued or running. When RQ cannot be asked,
    assume it is running: a resume must never start a second load."""
    try:
        from frappe.utils.background_jobs import is_job_enqueued
        return bool(is_job_enqueued(_job_id(name)))
    except Exception:
        return True


def _require_loader():
    if not (frappe.has_permission("Trial Balance Submission", "submit")
            and frappe.has_permission(DOCTYPE, "create")):
        frappe.throw("Loading trial balances in bulk needs the Close Lead role (EPM Admin).",
                     frappe.PermissionError)


# ── permissions (registered in hooks.py) ──────────────────────────────────

def upload_conditions(user=None):
    """List filter: a user limited to some entities sees only their own uploads."""
    user = user or frappe.session.user
    if allowed_entity_codes(user) is None:
        return ""
    return f"`tab{DOCTYPE}`.`owner` = {frappe.db.escape(user)}"


def has_upload_permission(doc, user=None, permission_type=None):
    """An upload holds every entity's figures (its report, and its file, whose
    read permission follows the upload). Only its owner, or a user who may
    see every entity in it, may open it."""
    user = user or frappe.session.user
    allowed = allowed_entity_codes(user)
    if allowed is None or doc.owner == user:
        return True
    entities = {r.get("entity") for r in json.loads(doc.report or "[]")}
    return entities <= set(allowed)


# ── reading and checking the file ─────────────────────────────────────────

def _xlsx_rows(content):
    from openpyxl import load_workbook
    if isinstance(content, str):
        content = content.encode("latin-1")
    wb = load_workbook(filename=BytesIO(content), data_only=True, read_only=True)
    try:
        # The first sheet, as the page promises; not whichever sheet happened
        # to be selected when the workbook was saved.
        return [list(row) for row in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def _read_table(file_url):
    file_doc = frappe.get_doc("File", {"file_url": file_url})
    if not frappe.has_permission("File", "read", doc=file_doc):
        frappe.throw("You cannot read that file.", frappe.PermissionError)
    name = (file_doc.file_name or file_url).lower()
    content = file_doc.get_content()
    if name.endswith(".xlsx"):
        return _xlsx_rows(content)
    if name.endswith((".csv", ".txt")):
        # Excel's "CSV UTF-8" starts with a byte-order mark; File.get_content
        # may already have decoded it into a string, so strip it either way.
        content = content.decode("utf-8-sig") if isinstance(content, bytes) else content.lstrip("﻿")
        return M.table_from_csv(content)
    frappe.throw("Upload a .csv or .xlsx file.")


def _chart_accounts():
    """The group chart, strictly: a warehouse outage refuses the check rather
    than passing unverified accounts (the single-upload rule)."""
    try:
        text = execute("SELECT DISTINCT main_account_id FROM epm_silver.silver_main_accounts")
    except Exception as e:
        frappe.throw(f"Cannot check accounts against the group chart: ClickHouse is unreachable ({e}). "
                     "Try again once the warehouse is up.")
    return {line.strip() for line in text.splitlines() if line.strip()}


def _check(table):
    groups = M.split_table(table)
    entities = sorted({k[0] for k in groups})
    # get_list applies the uploader's entity scope; get_all would not.
    visible = set(frappe.get_list("Entity", filters={"name": ["in", entities]}, pluck="name",
                                  limit_page_length=0))
    leaf = set(frappe.get_all("Entity", filters={"name": ["in", entities], "is_group": 0}, pluck="name"))
    statuses = {(y, p): period_status.get_status(y, p) for (_, y, p) in groups if 1 <= p <= 12}
    existing = {}
    for r in frappe.get_all("Trial Balance Submission",
                            filters={"docstatus": 1, "data_area_id": ["in", entities]},
                            fields=["name", "data_area_id", "fiscal_year", "fiscal_period"],
                            limit_page_length=0):
        existing[(r.data_area_id, int(r.fiscal_year), int(r.fiscal_period))] = r.name
    chart = _chart_accounts()
    report = [
        M.check_group(key, rows, known_accounts=chart, visible=key[0] in visible, leaf=key[0] in leaf,
                      period_status=statuses.get((key[1], key[2])), existing=existing.get(key),
                      validate_rows=validate_tb_rows)
        for key, rows in groups.items()
    ]
    return groups, report


def _record_check(doc):
    """Check the file and write the report onto the upload. Rows an earlier
    run of this upload loaded are carried forward as loaded."""
    previous = json.loads(doc.report or "[]")
    try:
        _, report = _check(_read_table(doc.upload_file))
    except ValueError as e:
        doc.status, doc.error, doc.report = "Failed", str(e), "[]"
        doc.group_count = doc.valid_count = doc.total_rows = doc.loaded_count = 0
    else:
        report = M.merge_loaded(report, previous)
        doc.report = json.dumps(report)
        doc.group_count = len(report)
        doc.valid_count = sum(1 for r in report if r["ok"])
        doc.loaded_count = sum(1 for r in report if r.get("loaded"))
        doc.total_rows = sum(r["rows"] for r in report)
        doc.status, doc.error = "Checked", ""
    doc.failed_count = 0
    doc.save()


def _payload(doc, refused=None):
    return {
        "name": doc.name, "status": doc.status, "file_url": doc.upload_file,
        "file_name": (doc.upload_file or "").rsplit("/", 1)[-1],
        "group_count": doc.group_count, "valid_count": doc.valid_count, "total_rows": doc.total_rows,
        "loaded_count": doc.loaded_count, "failed_count": doc.failed_count, "error": doc.error,
        "report": json.loads(doc.report or "[]"), "owner": frappe.utils.get_fullname(doc.owner),
        "creation": str(doc.creation),
        # Loading, but no job queued or running: the worker stopped.
        "stalled": doc.status == "Loading" and not _job_running(doc.name),
        "refused": refused,
    }


# ── endpoints ──────────────────────────────────────────────────────────────

@frappe.whitelist(methods=["POST"])
def check_file(file_url):
    """Create an upload for an uploaded file and report on it. Loads nothing."""
    _require_loader()
    doc = frappe.get_doc({"doctype": DOCTYPE, "upload_file": file_url, "status": "Draft"}).insert()
    _record_check(doc)
    return _payload(doc)


@frappe.whitelist(methods=["POST"])
def load(name, skip_invalid=0):
    """Re-check, then queue the load of every ready entity-period.

    Also resumes an upload whose load stopped or finished partly. New
    problems found by the re-check come back as `refused` with the fresh
    report (saved), so the page can show them and offer to skip them.
    """
    _require_loader()
    doc = frappe.get_doc(DOCTYPE, name)
    doc.check_permission("write")
    resumable = doc.status in ("Partly Loaded", "Failed") or (doc.status == "Loading" and not _job_running(doc.name))
    if doc.status != "Checked" and not resumable:
        frappe.throw("This upload is still loading." if doc.status == "Loading"
                     else f"This upload is {doc.status.lower()}; there is nothing left to load.")
    _record_check(doc)
    if doc.status == "Failed":
        return _payload(doc, refused="The file could not be read.")
    report = json.loads(doc.report or "[]")
    ready = sum(1 for r in report if r["ok"] and not r.get("loaded"))
    problems = sum(1 for r in report if not r["ok"])
    if not ready:
        return _payload(doc, refused="Nothing in this file is left to load.")
    if problems and not cint(skip_invalid):
        return _payload(doc, refused=f"{problems} of {len(report)} entity-periods have problems now. "
                                     f"Load only the {ready} that {'is' if ready == 1 else 'are'} ready, or fix the file.")
    doc.status = "Loading"
    doc.save()
    frappe.enqueue("konsol.tb_bulk.run_load", queue="long", timeout=3600, job_id=_job_id(doc.name),
                   deduplicate=True, enqueue_after_commit=True, upload=doc.name)
    return _payload(doc)


@frappe.whitelist(methods=["GET"])
def get_upload(name):
    doc = frappe.get_doc(DOCTYPE, name)
    doc.check_permission("read")
    return _payload(doc)


@frappe.whitelist(methods=["GET"])
def recent_uploads(limit=8):
    return frappe.get_list(DOCTYPE, fields=["name", "status", "upload_file", "group_count", "valid_count",
                                            "loaded_count", "failed_count", "owner", "creation"],
                           order_by="creation desc", limit_page_length=min(cint(limit) or 8, 50))


# ── the load job ───────────────────────────────────────────────────────────

def _save_progress(name, report, loaded, failed, status=None, error=None):
    values = {"report": json.dumps(report), "loaded_count": loaded, "failed_count": failed}
    if status:
        values["status"] = status
    if error is not None:
        values["error"] = error
    frappe.db.set_value(DOCTYPE, name, values, update_modified=bool(status))
    frappe.db.commit()


def _unclaim(batch_id):
    """A submission whose Frappe commit failed after on_submit claimed its
    batch must not stay counted in the warehouse: remove the claim."""
    if not batch_id:
        return
    try:
        execute(f"ALTER TABLE {CONTROL_TABLE} DELETE WHERE batch_id = '{_sql_str(batch_id)}' "
                "SETTINGS mutations_sync = 1")
    except Exception:
        frappe.log_error(title=f"Could not remove the claim for trial balance batch {batch_id}")


def run_load(upload):
    """Background job: one Trial Balance Submission per ready entity-period.

    Runs as the user who asked (Frappe sets the job's user), so permissions
    and entity scope apply to every row. Each entity-period commits on its
    own; a failure is recorded on that row and the rest carry on. A job
    timeout is not one row's failure: it stops the load, which can then be
    resumed.
    """
    from rq.timeouts import BaseTimeoutException

    doc = frappe.get_doc(DOCTYPE, upload)
    report = json.loads(doc.report or "[]")
    loaded = sum(1 for r in report if r.get("loaded"))
    failed = 0
    in_flight = None
    try:
        groups = M.split_table(_read_table(doc.upload_file))
        ready = [r for r in report if r["ok"] and not r.get("loaded")]
        total = loaded + len(ready)
        for i, item in enumerate(ready, start=1):
            key = (item["entity"], int(item["fiscal_year"]), int(item["fiscal_period"]))
            in_flight = None
            try:
                file_doc = frappe.get_doc({
                    "doctype": "File", "is_private": 1,
                    "file_name": f"{doc.name}-{key[0]}-{key[1]}-P{key[2]:02d}.csv",
                    "content": M.group_csv(groups[key]),
                }).insert()
                tbs = frappe.get_doc({"doctype": "Trial Balance Submission", "data_area_id": key[0],
                                      "fiscal_year": key[1], "fiscal_period": key[2],
                                      "tb_file": file_doc.file_url}).insert()
                # Attach before submitting: on_submit lands and claims the
                # rows, so nothing but the commit may follow it.
                frappe.db.set_value("File", file_doc.name,
                                    {"attached_to_doctype": "Trial Balance Submission",
                                     "attached_to_name": tbs.name}, update_modified=False)
                in_flight = tbs.batch_id
                tbs.submit()
                frappe.db.commit()
                in_flight = None
                item["loaded"] = tbs.name
                item.pop("load_error", None)
                loaded += 1
            except BaseTimeoutException:
                raise
            except Exception as e:
                frappe.db.rollback()
                _unclaim(in_flight)
                in_flight = None
                frappe.clear_messages()
                item["load_error"] = strip_html(str(e)) or type(e).__name__
                failed += 1
            if i % 10 == 0:
                _save_progress(doc.name, report, loaded, failed)
        _save_progress(doc.name, report, loaded, failed, status=M.outcome(loaded, failed, total))
    except Exception as e:
        frappe.db.rollback()
        _unclaim(in_flight)
        if not isinstance(e, BaseTimeoutException):
            frappe.log_error(title=f"Trial balance upload {doc.name} failed")
        reason = ("The load ran out of time; resume it to load the rest." if isinstance(e, BaseTimeoutException)
                  else strip_html(str(e)) or type(e).__name__)
        _save_progress(doc.name, report, loaded, failed, status="Partly Loaded" if loaded else "Failed",
                       error=reason)
